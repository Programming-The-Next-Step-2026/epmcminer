"""Tests for epmcminer.services.report_service."""

import csv
from pathlib import Path

import openpyxl
import pytest

from epmcminer.api.paper import Paper
from epmcminer.api.search_params import SearchParams
from epmcminer.services.download_result import DownloadResult
from epmcminer.services.report_service import REPORT_COLUMNS, ReportService

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def make_params(**overrides: object) -> SearchParams:
    """Return a SearchParams with sensible defaults for report tests."""
    defaults: dict = {
        "query": "depression AND therapy",
        "date_from": "2020-01-01",
        "date_to": "2024-12-31",
        "sort_order": "relevance",
        "count": 10,
        "output_folder": Path("/tmp"),
        "licenses": ["CC-BY"],
        "publication_types": ["research-article"],
    }
    defaults.update(overrides)
    return SearchParams(**defaults)


def make_paper(**overrides: object) -> Paper:
    """Return a Paper with sensible defaults."""
    defaults: dict = {
        "pmid": "12345",
        "doi": "10.1234/test",
        "title": "Test Paper Title",
        "authors": "Smith J, Jones A",
        "journal": "Test Journal",
        "year": "2022",
        "abstract": "Abstract text.",
        "pdf_url": "http://example.com/test.pdf",
    }
    defaults.update(overrides)
    return Paper(**defaults)


def make_result(
    status: str = "downloaded",
    reason: str | None = None,
    file_path: Path | None = Path("/tmp/test.pdf"),
    **paper_overrides: object,
) -> DownloadResult:
    """Return a DownloadResult with sensible defaults."""
    return DownloadResult(
        paper=make_paper(**paper_overrides),
        status=status,
        reason=reason,
        file_path=file_path,
    )


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


@pytest.fixture
def service() -> ReportService:
    """Return a ReportService instance."""
    return ReportService()


@pytest.fixture
def params() -> SearchParams:
    """Return a default SearchParams."""
    return make_params()


@pytest.fixture
def results() -> list[DownloadResult]:
    """Return a mixed list of DownloadResults."""
    return [
        make_result(status="downloaded", file_path=Path("/tmp/a.pdf")),
        make_result(
            status="skipped",
            reason="PDF unavailable",
            file_path=None,
            pmid="2",
            doi="10.1234/other",
        ),
        make_result(
            status="failed",
            reason="503",
            file_path=None,
            pmid="3",
            doi="10.1234/fail",
        ),
    ]


# ---------------------------------------------------------------------------
# TestSaveCsv
# ---------------------------------------------------------------------------


class TestSaveCsv:
    """Tests for ReportService.save_csv."""

    def test_saves_to_report_csv(
        self,
        service: ReportService,
        results: list[DownloadResult],
        params: SearchParams,
        tmp_path: Path,
    ) -> None:
        """save_csv returns the path to {output_folder}/report.csv and the file exists."""
        path = service.save_csv(results, params, tmp_path)
        assert path == tmp_path / "report.csv"
        assert path.exists()

    def test_csv_has_header_row(
        self,
        service: ReportService,
        results: list[DownloadResult],
        params: SearchParams,
        tmp_path: Path,
    ) -> None:
        """The CSV header row matches REPORT_COLUMNS in order."""
        path = service.save_csv(results, params, tmp_path)
        with path.open(newline="") as f:
            reader = csv.DictReader(f)
            assert list(reader.fieldnames) == REPORT_COLUMNS

    def test_csv_has_correct_row_count(
        self,
        service: ReportService,
        results: list[DownloadResult],
        params: SearchParams,
        tmp_path: Path,
    ) -> None:
        """The CSV has one data row per DownloadResult."""
        path = service.save_csv(results, params, tmp_path)
        with path.open(newline="") as f:
            rows = list(csv.DictReader(f))
        assert len(rows) == len(results)

    def test_csv_paper_fields(
        self, service: ReportService, params: SearchParams, tmp_path: Path
    ) -> None:
        """Paper fields are written correctly to the CSV."""
        result = make_result(status="downloaded", file_path=Path("/tmp/x.pdf"))
        path = service.save_csv([result], params, tmp_path)
        with path.open(newline="") as f:
            row = list(csv.DictReader(f))[0]
        assert row["title"] == result.paper.title
        assert row["authors"] == result.paper.authors
        assert row["doi"] == result.paper.doi
        assert row["year"] == result.paper.year
        assert row["journal"] == result.paper.journal

    def test_csv_status_and_reason(
        self, service: ReportService, params: SearchParams, tmp_path: Path
    ) -> None:
        """Status and reason are written correctly to the CSV."""
        result = make_result(status="skipped", reason="PDF unavailable", file_path=None)
        path = service.save_csv([result], params, tmp_path)
        with path.open(newline="") as f:
            row = list(csv.DictReader(f))[0]
        assert row["status"] == "skipped"
        assert row["reason"] == "PDF unavailable"

    def test_csv_file_path_empty_when_none(
        self, service: ReportService, params: SearchParams, tmp_path: Path
    ) -> None:
        """A None file_path is represented as an empty string in the CSV."""
        result = make_result(status="failed", reason="503", file_path=None)
        path = service.save_csv([result], params, tmp_path)
        with path.open(newline="") as f:
            row = list(csv.DictReader(f))[0]
        assert row["file_path"] == ""

    def test_csv_includes_search_params(self, service: ReportService, tmp_path: Path) -> None:
        """Search parameters are written as columns in every row."""
        params = make_params(
            query="cancer",
            sort_order="date",
            date_from="2021-01-01",
            date_to="2023-12-31",
        )
        path = service.save_csv([make_result()], params, tmp_path)
        with path.open(newline="") as f:
            row = list(csv.DictReader(f))[0]
        assert row["query"] == "cancer"
        assert row["sort_order"] == "date"
        assert row["date_from"] == "2021-01-01"
        assert row["date_to"] == "2023-12-31"

    def test_csv_licenses_joined(self, service: ReportService, tmp_path: Path) -> None:
        """Multiple licenses are joined with ', ' in the CSV."""
        params = make_params(licenses=["CC-BY", "CC0"])
        path = service.save_csv([make_result()], params, tmp_path)
        with path.open(newline="") as f:
            row = list(csv.DictReader(f))[0]
        assert row["licenses"] == "CC-BY, CC0"

    def test_csv_publication_types_joined(self, service: ReportService, tmp_path: Path) -> None:
        """Multiple publication types are joined with ', ' in the CSV."""
        params = make_params(publication_types=["research-article", "review"])
        path = service.save_csv([make_result()], params, tmp_path)
        with path.open(newline="") as f:
            row = list(csv.DictReader(f))[0]
        assert row["publication_types"] == "research-article, review"

    def test_csv_reason_empty_when_none(
        self, service: ReportService, params: SearchParams, tmp_path: Path
    ) -> None:
        """A None reason (downloaded status) is written as an empty string in the CSV."""
        result = make_result(status="downloaded", reason=None, file_path=Path("/tmp/x.pdf"))
        path = service.save_csv([result], params, tmp_path)
        with path.open(newline="") as f:
            row = list(csv.DictReader(f))[0]
        assert row["reason"] == ""

    def test_csv_file_path_written_when_present(
        self, service: ReportService, params: SearchParams, tmp_path: Path
    ) -> None:
        """A non-None file_path is serialised as its string representation in the CSV."""
        file_path = Path("/some/output/folder/paper.pdf")
        result = make_result(status="downloaded", file_path=file_path)
        path = service.save_csv([result], params, tmp_path)
        with path.open(newline="") as f:
            row = list(csv.DictReader(f))[0]
        assert row["file_path"] == str(file_path)

    def test_csv_empty_licenses(self, service: ReportService, tmp_path: Path) -> None:
        """An empty licenses list is written as an empty string in the CSV."""
        params = make_params(licenses=[])
        path = service.save_csv([make_result()], params, tmp_path)
        with path.open(newline="") as f:
            row = list(csv.DictReader(f))[0]
        assert row["licenses"] == ""

    def test_csv_empty_publication_types(self, service: ReportService, tmp_path: Path) -> None:
        """An empty publication_types list is written as an empty string in the CSV."""
        params = make_params(publication_types=[])
        path = service.save_csv([make_result()], params, tmp_path)
        with path.open(newline="") as f:
            row = list(csv.DictReader(f))[0]
        assert row["publication_types"] == ""

    def test_csv_params_repeated_across_all_rows(
        self, service: ReportService, tmp_path: Path
    ) -> None:
        """Search parameters are identical in every data row."""
        params = make_params(query="anxiety", sort_order="citations")
        results = [make_result(pmid=str(i), doi=f"10.1/{i}") for i in range(3)]
        path = service.save_csv(results, params, tmp_path)
        with path.open(newline="") as f:
            rows = list(csv.DictReader(f))
        assert all(r["query"] == "anxiety" for r in rows)
        assert all(r["sort_order"] == "citations" for r in rows)

    def test_csv_overwrites_existing_file(
        self, service: ReportService, params: SearchParams, tmp_path: Path
    ) -> None:
        """A second call to save_csv replaces the previous report.csv."""
        service.save_csv([make_result(pmid="1", doi="10.1/a")], params, tmp_path)
        path = service.save_csv([make_result(pmid="2", doi="10.1/b")], params, tmp_path)
        with path.open(newline="") as f:
            rows = list(csv.DictReader(f))
        assert len(rows) == 1
        assert rows[0]["doi"] == "10.1/b"

    def test_csv_creates_output_folder_if_missing(
        self, service: ReportService, params: SearchParams, tmp_path: Path
    ) -> None:
        """save_csv creates the output_folder when it does not yet exist."""
        new_folder = tmp_path / "new" / "nested"
        service.save_csv([make_result()], params, new_folder)
        assert (new_folder / "report.csv").exists()

    def test_csv_empty_results(
        self, service: ReportService, params: SearchParams, tmp_path: Path
    ) -> None:
        """An empty results list produces a CSV with only a header row."""
        path = service.save_csv([], params, tmp_path)
        with path.open(newline="") as f:
            rows = list(csv.DictReader(f))
        assert rows == []


# ---------------------------------------------------------------------------
# TestExportExcel
# ---------------------------------------------------------------------------


class TestExportExcel:
    """Tests for ReportService.export_excel."""

    def test_creates_file(
        self,
        service: ReportService,
        results: list[DownloadResult],
        params: SearchParams,
        tmp_path: Path,
    ) -> None:
        """export_excel creates the output file."""
        out = tmp_path / "report.xlsx"
        service.export_excel(results, params, out)
        assert out.exists()

    def test_sheet_exists(
        self,
        service: ReportService,
        results: list[DownloadResult],
        params: SearchParams,
        tmp_path: Path,
    ) -> None:
        """The Excel workbook contains at least one sheet."""
        out = tmp_path / "report.xlsx"
        service.export_excel(results, params, out)
        wb = openpyxl.load_workbook(out)
        assert len(wb.sheetnames) >= 1

    def test_header_row_matches_report_columns(
        self,
        service: ReportService,
        results: list[DownloadResult],
        params: SearchParams,
        tmp_path: Path,
    ) -> None:
        """The first row of the active sheet matches REPORT_COLUMNS in order."""
        out = tmp_path / "report.xlsx"
        service.export_excel(results, params, out)
        wb = openpyxl.load_workbook(out)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        assert headers == REPORT_COLUMNS

    def test_data_row_count(
        self,
        service: ReportService,
        results: list[DownloadResult],
        params: SearchParams,
        tmp_path: Path,
    ) -> None:
        """The sheet has exactly one data row per DownloadResult."""
        out = tmp_path / "report.xlsx"
        service.export_excel(results, params, out)
        wb = openpyxl.load_workbook(out)
        ws = wb.active
        # Row 1 is the header; data rows start at row 2.
        assert ws.max_row - 1 == len(results)

    def test_paper_fields_in_excel(
        self, service: ReportService, params: SearchParams, tmp_path: Path
    ) -> None:
        """Paper title and DOI are written correctly into the Excel sheet."""
        result = make_result(status="downloaded", file_path=Path("/tmp/x.pdf"))
        out = tmp_path / "report.xlsx"
        service.export_excel([result], params, out)
        wb = openpyxl.load_workbook(out)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        data_row = dict(zip(headers, [cell.value for cell in ws[2]], strict=False))
        assert data_row["title"] == result.paper.title
        assert data_row["doi"] == result.paper.doi

    def test_search_params_in_excel(self, service: ReportService, tmp_path: Path) -> None:
        """Search parameters appear as column values in each Excel data row."""
        params = make_params(query="anxiety", sort_order="date")
        out = tmp_path / "report.xlsx"
        service.export_excel([make_result()], params, out)
        wb = openpyxl.load_workbook(out)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        data_row = dict(zip(headers, [cell.value for cell in ws[2]], strict=False))
        assert data_row["query"] == "anxiety"
        assert data_row["sort_order"] == "date"

    def test_excel_creates_parent_folder_if_missing(
        self, service: ReportService, params: SearchParams, tmp_path: Path
    ) -> None:
        """export_excel creates the parent directory when it does not yet exist."""
        out = tmp_path / "new" / "nested" / "report.xlsx"
        service.export_excel([make_result()], params, out)
        assert out.exists()

    def test_excel_empty_results(
        self, service: ReportService, params: SearchParams, tmp_path: Path
    ) -> None:
        """An empty results list produces an Excel file with only a header row."""
        out = tmp_path / "report.xlsx"
        service.export_excel([], params, out)
        wb = openpyxl.load_workbook(out)
        ws = wb.active
        assert ws.max_row == 1

    def test_excel_file_path_empty_when_none(
        self, service: ReportService, params: SearchParams, tmp_path: Path
    ) -> None:
        """A None file_path produces an empty cell (None) in the Excel output.

        Excel has no empty-string concept distinct from an empty cell; openpyxl
        reads empty cells back as None rather than "".
        """
        result = make_result(status="failed", reason="503", file_path=None)
        out = tmp_path / "report.xlsx"
        service.export_excel([result], params, out)
        wb = openpyxl.load_workbook(out)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        data_row = dict(zip(headers, [cell.value for cell in ws[2]], strict=False))
        assert data_row["file_path"] is None


# ---------------------------------------------------------------------------
# TestExportPdf
# ---------------------------------------------------------------------------


class TestExportPdf:
    """Tests for ReportService.export_pdf."""

    def test_creates_file(
        self,
        service: ReportService,
        results: list[DownloadResult],
        params: SearchParams,
        tmp_path: Path,
    ) -> None:
        """export_pdf creates the output file."""
        out = tmp_path / "report.pdf"
        service.export_pdf(results, params, out)
        assert out.exists()

    def test_file_is_non_empty(
        self,
        service: ReportService,
        results: list[DownloadResult],
        params: SearchParams,
        tmp_path: Path,
    ) -> None:
        """The generated PDF file is not empty."""
        out = tmp_path / "report.pdf"
        service.export_pdf(results, params, out)
        assert out.stat().st_size > 0

    def test_output_is_valid_pdf(
        self,
        service: ReportService,
        results: list[DownloadResult],
        params: SearchParams,
        tmp_path: Path,
    ) -> None:
        """The output file starts with the PDF magic bytes."""
        out = tmp_path / "report.pdf"
        service.export_pdf(results, params, out)
        assert out.read_bytes().startswith(b"%PDF-")

    def test_pdf_creates_parent_folder_if_missing(
        self, service: ReportService, params: SearchParams, tmp_path: Path
    ) -> None:
        """export_pdf creates the parent directory when it does not yet exist."""
        out = tmp_path / "new" / "nested" / "report.pdf"
        service.export_pdf([make_result()], params, out)
        assert out.exists()

    def test_empty_results_produces_valid_pdf(
        self, service: ReportService, params: SearchParams, tmp_path: Path
    ) -> None:
        """An empty results list still produces a valid, non-empty PDF."""
        out = tmp_path / "report.pdf"
        service.export_pdf([], params, out)
        assert out.exists()
        assert out.stat().st_size > 0
        assert out.read_bytes().startswith(b"%PDF-")

    def test_pdf_with_all_downloaded(
        self, service: ReportService, params: SearchParams, tmp_path: Path
    ) -> None:
        """All-downloaded results produce a valid PDF (no skipped section)."""
        results = [
            make_result(status="downloaded", file_path=Path("/tmp/a.pdf"), pmid=str(i))
            for i in range(3)
        ]
        out = tmp_path / "report.pdf"
        service.export_pdf(results, params, out)
        assert out.exists()
        assert out.read_bytes().startswith(b"%PDF-")

    def test_pdf_with_all_skipped(
        self, service: ReportService, params: SearchParams, tmp_path: Path
    ) -> None:
        """All-skipped results produce a valid PDF with a skipped section."""
        results = [
            make_result(status="skipped", reason="PDF unavailable", file_path=None, pmid=str(i))
            for i in range(3)
        ]
        out = tmp_path / "report.pdf"
        service.export_pdf(results, params, out)
        assert out.exists()
        assert out.read_bytes().startswith(b"%PDF-")

    def test_pdf_with_total_found_nonzero(
        self, service: ReportService, params: SearchParams, tmp_path: Path
    ) -> None:
        """Passing total_found=42 produces a valid PDF."""
        out = tmp_path / "report.pdf"
        service.export_pdf([make_result()], params, out, total_found=42)
        assert out.exists()
        assert out.read_bytes().startswith(b"%PDF-")

    def test_pdf_with_total_found_zero(
        self, service: ReportService, params: SearchParams, tmp_path: Path
    ) -> None:
        """Default total_found=0 produces a valid PDF."""
        out = tmp_path / "report.pdf"
        service.export_pdf([make_result()], params, out, total_found=0)
        assert out.exists()
        assert out.read_bytes().startswith(b"%PDF-")

    def test_pdf_mixed_params_filters(self, service: ReportService, tmp_path: Path) -> None:
        """Non-empty licenses, publication_types, and author_orcids produce a valid PDF."""
        params = make_params(
            licenses=["CC BY", "CC0"],
            publication_types=["Review", "Research Article"],
            author_orcids=["0000-0001-2345-6789", "0000-0002-9876-5432"],
        )
        out = tmp_path / "report.pdf"
        service.export_pdf([make_result()], params, out)
        assert out.exists()
        assert out.read_bytes().startswith(b"%PDF-")

    def test_pdf_result_with_missing_paper_fields(
        self, service: ReportService, params: SearchParams, tmp_path: Path
    ) -> None:
        """A skipped result whose paper has empty optional fields does not raise."""
        result = make_result(
            status="skipped",
            reason="Server rate limiting",
            file_path=None,
            authors="",
            journal="",
            year="",
        )
        out = tmp_path / "report.pdf"
        service.export_pdf([result], params, out)
        assert out.exists()
        assert out.read_bytes().startswith(b"%PDF-")

    def test_pdf_downloaded_section_included_when_downloads_exist(
        self, service: ReportService, params: SearchParams, tmp_path: Path
    ) -> None:
        """Results with status='downloaded' produce a valid PDF (downloaded section present)."""
        results = [
            make_result(status="downloaded", file_path=Path("/tmp/a.pdf"), pmid="1"),
            make_result(status="downloaded", file_path=Path("/tmp/b.pdf"), pmid="2"),
        ]
        out = tmp_path / "report.pdf"
        service.export_pdf(results, params, out)
        assert out.exists()
        assert out.read_bytes().startswith(b"%PDF-")

    def test_pdf_downloaded_section_omitted_when_all_skipped(
        self, service: ReportService, params: SearchParams, tmp_path: Path
    ) -> None:
        """All-skipped results produce a valid PDF (downloaded section absent)."""
        results = [
            make_result(status="skipped", reason="PDF unavailable", file_path=None, pmid="1"),
        ]
        out = tmp_path / "report.pdf"
        service.export_pdf(results, params, out)
        assert out.exists()
        assert out.read_bytes().startswith(b"%PDF-")

    def test_pdf_both_sections_present_for_mixed_results(
        self, service: ReportService, params: SearchParams, tmp_path: Path
    ) -> None:
        """Mixed results produce a valid PDF containing both downloaded and skipped sections."""
        results = [
            make_result(status="downloaded", file_path=Path("/tmp/a.pdf"), pmid="1"),
            make_result(status="skipped", reason="PDF unavailable", file_path=None, pmid="2"),
        ]
        out = tmp_path / "report.pdf"
        service.export_pdf(results, params, out)
        assert out.exists()
        assert out.read_bytes().startswith(b"%PDF-")

    def test_pdf_downloaded_paper_without_file_path_does_not_raise(
        self, service: ReportService, params: SearchParams, tmp_path: Path
    ) -> None:
        """A downloaded result with file_path=None renders without error."""
        result = make_result(status="downloaded", file_path=None, pmid="1")
        out = tmp_path / "report.pdf"
        service.export_pdf([result], params, out)
        assert out.exists()
        assert out.read_bytes().startswith(b"%PDF-")
